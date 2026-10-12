"""Read a design table: one Project No. block of a project workbook, or a CSV file.

The result is a list of ``Row`` objects. Each row holds the values of one line of
the table by column name (the names in ``config.json`` -> ``columns``), the line
number it came from, and the Sub. No. it belongs to. A line whose Sub. No. is
empty is a continuation of the sub above it (the second PCR of a construct).

This module only reads. Writing the protocol workbook is ``make_protocol.py``.
"""

from __future__ import annotations

import csv
import datetime as dt
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException


class DesignError(Exception):
    """The design table cannot be used; the message says why."""


@dataclass
class Row:
    line: int                      # row number in the workbook, or line number in the CSV
    values: dict[str, object]      # column name -> value
    sub: str = ""                  # Sub. No. of this row, inherited by continuation rows
    own: bool = False              # True if this row carries its own Sub. No.

    def text(self, name: str) -> str:
        return cell_text(self.values.get(name))


@dataclass
class Block:
    header: int
    first: int
    last: int
    project: str | None
    label: str
    rows: list[Row] = field(default_factory=list)
    headings: dict[int, str] = field(default_factory=dict)   # column number -> heading text
    unsaved: list[str] = field(default_factory=list)         # formula cells with no saved value
    stray: list[str] = field(default_factory=list)           # notes about rows not read as data


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y/%m/%d")
    return str(value).strip()


def norm_sub(value: object) -> str:
    """``3``, ``3.0`` and ``"03"`` are all sub 3; anything else is kept as text."""
    text = cell_text(value)
    number = to_number(text)
    if isinstance(number, int):
        return str(number)
    return text


def to_number(text: str) -> int | float | None:
    """The number a CSV cell reads as, or None. ``"01"`` is 1."""
    s = text.strip()
    if not s:
        return None
    try:
        return int(s)
    except ValueError:
        pass
    try:
        number = float(s)
    except ValueError:
        return None
    if number != number or number in (float("inf"), float("-inf")):
        return None
    return int(number) if number.is_integer() else number


def _clean(value: object) -> object:
    """A workbook cell as a value: integral floats become ints, text is trimmed."""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.strip() or None
    return value


def assign_subs(rows: list[Row]) -> None:
    current = ""
    for row in rows:
        sub = norm_sub(row.values.get("sub_no"))
        if sub:
            current = sub
        row.sub = current
        row.own = bool(sub)


# --- xlsx ----------------------------------------------------------------------------------

def read_blocks(path: str | Path, sheet: str, cfg: dict) -> list[Block]:
    """Every block of ``sheet``: a header row (B = ``Project No.``) and the rows under it."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except (zipfile.BadZipFile, InvalidFileException, OSError, KeyError) as exc:
        raise DesignError(f"cannot read {path} as a workbook: {exc}") from None
    if sheet not in wb.sheetnames:
        raise DesignError(
            f"sheet '{sheet}' not found in {path} (sheets: {', '.join(wb.sheetnames)})")
    names = {idx: name for name, idx in cfg["columns"].items()}
    width = max(names) + 1
    ws = wb[sheet]
    ws.reset_dimensions()          # some writers store a wrong used range
    grid = [list(r) + [None] * (width - len(r)) for r in
            ws.iter_rows(min_col=1, max_col=width, values_only=True)]
    wb.close()
    raw = openpyxl.load_workbook(path, read_only=True)[sheet]       # formulas, not values
    raw.reset_dimensions()
    formulas = {(i, c) for i, r in enumerate(raw.iter_rows(min_col=1, max_col=width,
                                                          values_only=True), start=1)
                for c, v in enumerate(r) if isinstance(v, str) and v.startswith("=")}

    def empty(r: int) -> bool:
        return all(cell_text(v) == "" for v in grid[r - 1])

    def is_header(r: int) -> bool:
        return cell_text(grid[r - 1][1]) == "Project No."

    blocks: list[Block] = []
    r, last_row = 1, len(grid)
    while r <= last_row:
        if not is_header(r):
            r += 1
            continue
        end = r
        while end + 1 <= last_row and not empty(end + 1) and not is_header(end + 1):
            end += 1
        if end > r:
            first = cell_text(grid[r][1]).upper() or None
            block = Block(header=r, first=r + 1, last=end, project=first,
                          label=cell_text(grid[r - 1][0]),
                          headings={c + 1: cell_text(v) for c, v in enumerate(grid[r - 1])})
            formula_cells, saved = [], 0
            for line in range(r + 1, end + 1):
                cells = grid[line - 1]
                values = {name: _clean(cells[idx - 1]) for idx, name in names.items()}
                block.rows.append(Row(line=line, values=values))
                for idx, name in names.items():
                    if (line, idx - 1) in formulas:
                        saved += cells[idx - 1] is not None
                        if name != "project_no":          # the Project No. is filled in anyway
                            formula_cells.append(f"line {line} {name}")
            # A workbook saved by Excel stores a result for its formulas (an empty result is
            # the only one that reads as None); one written by a program stores none at all.
            if formula_cells and saved == 0:
                block.unsaved = formula_cells
            others = sorted({row.text("project_no").upper() for row in block.rows
                             if row.text("project_no")} - {first or ""})
            if first and others:
                block.stray.append(
                    f"the Project No. in column B changes inside the block ({', '.join(others)}); "
                    f"every row is written as {first}.")
            beyond = []
            nxt = end + 1
            while nxt <= last_row and not is_header(nxt):
                if not empty(nxt) and first and cell_text(grid[nxt - 1][1]).upper() == first:
                    beyond.append(str(nxt))
                nxt += 1
            if beyond:
                block.stray.append(
                    f"line(s) {', '.join(beyond[:8])}{' ...' if len(beyond) > 8 else ''} carry "
                    f"{first} but follow a blank row, so they are not read. Remove the blank row.")
            assign_subs(block.rows)
            blocks.append(block)
        r = end + 1
    return blocks


def pick_block(blocks: list[Block], project: str | None, label: str, sheet: str) -> Block:
    if not blocks:
        raise DesignError(
            f"no block found in sheet '{sheet}'. A block is a row whose column B reads "
            "'Project No.' with at least one row of data under it.")
    wanted = project.strip().upper() if project else None
    known = ", ".join(b.project or "(blank)" for b in blocks)
    if wanted:
        hits = [b for b in blocks if b.project == wanted]
        if not hits:
            blank = [b for b in blocks if b.project is None]
            if len(blank) == 1:
                hits = blank
        if not hits:
            raise DesignError(f"Project {project} not found in '{sheet}'. Blocks present: {known}")
        if len(hits) > 1:
            raise DesignError(
                f"Project {project} has several blocks (header rows "
                f"{', '.join(str(b.header) for b in hits)}). Fix the workbook first.")
        block = hits[0]
    else:
        if len(blocks) != 1:
            raise DesignError(f"'{sheet}' has {len(blocks)} blocks ({known}); pass --project.")
        block = blocks[0]
        if block.project is None:
            raise DesignError("the block has no Project No. in column B; pass --project.")
    if block.label != label:
        raise DesignError(
            f"Block {block.project or project} (row {block.header}) is '{block.label}', "
            f"not '{label}'; the plasmid-construction protocol does not apply.")
    return block


# --- CSV -----------------------------------------------------------------------------------

def read_csv(path: str | Path, cfg: dict) -> tuple[list[Row], list[str]]:
    """Rows of a CSV design table and the (normalised) header names it has.

    Headers are matched ignoring case and surrounding spaces; unknown columns are
    ignored. An empty ``project_no`` repeats the one above it. Values stay text
    here (empty -> None); ``coerce`` converts the numeric ones.
    """
    known = set(cfg["columns"])
    try:
        with open(path, encoding="utf-8-sig", newline="") as fh:
            records = list(csv.reader(fh))
    except UnicodeDecodeError:
        raise DesignError(
            f"{path} is not UTF-8 text. In Excel, save a table as 'CSV UTF-8 (comma "
            "delimited)'; a workbook is given with its .xlsx name.") from None
    except OSError as exc:
        raise DesignError(f"cannot read {path}: {exc.strerror or exc}. "
                          "If it is open in Excel, close it.") from None
    if not records:
        raise DesignError(f"{path} is empty")
    names = [h.strip().lower() for h in records[0]]
    if not any(n in known for n in names):
        raise DesignError(
            f"none of the headings in {path} is a column name. The headings must be names "
            "such as sub_no, plasmid_name, re1, re2, backbone, oligo1, oligo2 (see "
            "docs/design-table.md), separated by commas.")
    rows: list[Row] = []
    for line, record in enumerate(records[1:], start=2):
        if not any(c.strip() for c in record):
            continue
        values: dict[str, object] = {name: None for name in known}
        for name, cell in zip(names, record, strict=False):
            if name in known:
                values[name] = cell.strip() or None
        rows.append(Row(line=line, values=values))
    seen = [n for n in names if n in known]
    carry = None
    for row in rows:
        if row.values.get("project_no"):
            carry = row.values["project_no"]
        elif carry is not None:
            row.values["project_no"] = carry
    assign_subs(rows)
    return rows, seen


def coerce(rows: list[Row], numeric_required: list[str]) -> None:
    """Turn CSV text into numbers where the protocol sheets need numbers.

    ``numeric_required`` columns must read as numbers when they are not empty, or
    ``DesignError`` names the row and column. ``length`` and ``conc`` are free
    text that is turned into a number only when it reads as one.
    """
    for row in rows:
        for name in numeric_required:
            raw = row.values.get(name)
            if raw is None or not isinstance(raw, str):
                continue
            number = to_number(raw)
            if number is None:
                raise DesignError(
                    f"line {row.line}: '{name}' must be a number, got '{raw}'")
            row.values[name] = number
        for name in ("length", "conc"):
            raw = row.values.get(name)
            if isinstance(raw, str):
                number = to_number(raw)
                if number is not None:
                    row.values[name] = number
    assign_subs(rows)


# --- selection -----------------------------------------------------------------------------

def select_rows(rows: list[Row], wanted: list[str]) -> tuple[list[Row], list[str]]:
    """The rows for ``--sub``, each with its continuation rows. Returns (rows, notes)."""
    if not wanted:
        return list(rows), []
    notes: list[str] = []
    picked: list[int] = []
    for sub in wanted:
        starts = [i for i, r in enumerate(rows) if r.own and r.sub == sub]
        if not starts:
            present = ", ".join(dict.fromkeys(r.sub for r in rows if r.own)) or "none"
            raise DesignError(f"Sub. No. {sub} not found. Sub. Nos. present: {present}. "
                              "Give several as --sub 3,4,5.")
        start = starts[-1]
        if len(starts) > 1:
            notes.append(f"Sub {sub} appears {len(starts)} times (lines "
                         f"{', '.join(str(rows[i].line) for i in starts)}); "
                         f"using the last, line {rows[start].line}.")
        picked.append(start)
        for i in range(start + 1, len(rows)):
            if rows[i].own:
                break
            picked.append(i)
    shared = [r.line for r in rows if r.sub == ""]
    if shared:
        notes.append(f"lines {', '.join(str(n) for n in shared)} have no Sub. No. "
                     "(shared step) and are left out; drop --sub to include them.")
    return [rows[i] for i in sorted(set(picked))], notes


def split_subs(values: list[str] | None) -> list[str]:
    out: list[str] = []
    for item in values or []:
        for piece in re.split(r"[,\s]+", item):
            if piece:
                out.append(norm_sub(piece))
    if values is not None and not out:
        raise DesignError("--sub is empty. Give the sub numbers, for example --sub 3,4,5.")
    return out
