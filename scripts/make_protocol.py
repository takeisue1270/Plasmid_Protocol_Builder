"""Build a protocol workbook (HiFi or GGA) from a design table.

    python scripts/make_protocol.py --design examples/design_table.csv --date 261011
    python scripts/make_protocol.py --design projects.xlsx --project EX001 --sub 3,4,5

The design table is a CSV file, or one Project No. block of a project workbook
(see docs/design-table.md). The rows are written to the Data sheet of a copy of
the template; every other sheet is formulas that read Data. Excel computes them
when the workbook is opened.

Output: <out-dir>/YYMMDD.<project>.<HiFi|GGA>.TF.xlsx
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string
from openpyxl.utils.exceptions import IllegalCharacterError, InvalidFileException

sys.path.insert(0, str(Path(__file__).resolve().parent))

import design_table as dtab  # noqa: E402
from design_table import DesignError  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
COMMENT_SIZE = (300, 120)  # width, height in pixels
RELIST_RANGE = ("A4", "D650")


def load_config(path: Path | None = None) -> dict:
    where = path or ROOT / "config.json"
    try:
        with open(where, encoding="utf-8") as fh:
            cfg = json.load(fh)
    except (OSError, json.JSONDecodeError) as exc:
        raise DesignError(f"cannot read the configuration {where}: {exc}") from None
    keys = ("data_sheet", "design_sheet", "construction_label", "default_method", "columns",
            "methods", "header_check")
    if not isinstance(cfg, dict) or any(k not in cfg for k in keys):
        raise DesignError(f"{where} is not a configuration of this tool; it needs the keys "
                          f"{', '.join(keys)} (compare config.json in the repository).")
    return cfg


def resolve_method(cfg: dict, name: str) -> str:
    hits = [m for m in cfg["methods"] if m.lower() == name.strip().lower()]
    if len(hits) != 1:
        raise DesignError(f"method must be one of: {', '.join(cfg['methods'])} (got '{name}')")
    return hits[0]


def open_template(path: Path, cfg: dict, spec: dict, method: str):
    """Load a template and check that it is a workbook built for this method."""
    try:
        wb = openpyxl.load_workbook(path)
    except (zipfile.BadZipFile, InvalidFileException, OSError, KeyError) as exc:
        raise DesignError(f"cannot read the template {path} as a workbook: {exc}") from None
    for sheet in (cfg["data_sheet"], spec["required_sheet"]):
        if sheet not in wb.sheetnames:
            raise DesignError(
                f"the template {path} has no '{sheet}' sheet, which a {method} template needs "
                f"(its sheets: {', '.join(wb.sheetnames)}).")
    return wb


def read_relist(wb) -> tuple[dict[str, dict], list[str]] | None:
    """The template's RElist as (enzyme name in lower case -> row, warnings), or None.

    The RE sheet matches names exactly and reads columns C and D, so a name with a
    space at either end, and a row with an empty cell in B to D, are reported.
    """
    if "RElist" not in wb.sheetnames:
        return None
    ws = wb["RElist"]
    table: dict[str, dict] = {}
    warns: list[str] = []
    for a, b, c, d in ws[f"{RELIST_RANGE[0]}:{RELIST_RANGE[1]}"]:
        name = dtab.cell_text(a.value)
        if not name:
            continue
        if isinstance(a.value, str) and a.value != a.value.strip():
            warns.append(f"RElist row {a.row}: the name '{a.value}' has a space at an end, so "
                         "the RE sheet cannot find it. Remove the space.")
        empty = [x.column_letter for x in (b, c, d) if dtab.cell_text(x.value) == ""]
        if empty:
            warns.append(f"RElist row {a.row} ({name}): column(s) {', '.join(empty)} are empty. "
                         "The RE sheet needs all four cells.")
        table[name.lower()] = {"name": name, "buffer": dtab.cell_text(b.value),
                               "temp": dtab.cell_text(c.value), "heat": dtab.cell_text(d.value)}
    return table, warns


def check_relist(rows: list[dtab.Row], table: dict[str, dict]) -> list[str]:
    warns: list[str] = []
    for row in rows:
        where = f"line {row.line} ({row.text('plasmid_name') or 'sub ' + row.sub})"
        first = row.text("re1")
        if not first:
            continue          # the RE sheet leaves a row without re1 blank
        if table.get(first.lower()) is None:
            # The RE sheet marks the row PCR and leaves temperature and heat inactivation
            # blank, so RE2 is not looked at.
            warns.append(
                f"{where}: re1 '{first}' is not in the RElist sheet. The RE sheet treats it "
                "as a backbone prepared by PCR. If it is an enzyme, add it to RElist.")
            continue
        for col in ("re1", "re2"):
            enzyme = row.text(col)
            if not enzyme:
                continue
            hit = table.get(enzyme.lower())
            if hit is None:
                warns.append(
                    f"{where}: re2 '{enzyme}' is not in the RElist sheet. The RE sheet shows "
                    "'Check' for the reaction temperature (and for heat inactivation unless "
                    "column purification is Yes).")
            elif hit["buffer"] and hit["buffer"].lower() != "cutsmart":
                warns.append(
                    f"{where}: the RE sheet is built for 10×CutSmart. The buffer of "
                    f"'{enzyme}' is {hit['buffer']}.")
    return warns


def keep_method(rows: list[dtab.Row], method: str) -> tuple[list[dtab.Row], list[dtab.Row]]:
    """Split rows into those for ``method`` and those for another method.

    A row without a method, and a continuation row (which takes the method of the sub
    above it), stay with the method that was asked for.
    """
    kept, dropped, current = [], [], ""
    for row in rows:
        if row.own:
            current = row.text("method")
        (kept if not current or current.lower() == method.lower() else dropped).append(row)
    return kept, dropped


def structure_warnings(rows: list[dtab.Row]) -> list[str]:
    """Rows the sheets cannot place: before the first Sub. No., and a Sub. No. used twice."""
    warns = []
    leading = [str(r.line) for r in rows if r.sub == ""]
    if leading:
        warns.append(f"line(s) {', '.join(leading)} have no Sub. No. and no sub above them. "
                     "The sheets cannot place such a row (the HiFi sheet can show an error). "
                     "Give it a Sub. No., or leave it out with --sub.")
    seen: dict[str, list[str]] = {}
    for r in rows:
        if r.own:
            seen.setdefault(r.sub, []).append(str(r.line))
    for sub_no, lines in seen.items():
        if len(lines) > 1:
            warns.append(f"Sub. No. {sub_no} is used on lines {', '.join(lines)}. Each row "
                         "with a Sub. No. starts a new construct; a repeat is counted twice.")
    return warns


def header_warnings(block: dtab.Block, cfg: dict) -> list[str]:
    warns = []
    for letter, accepted in cfg["header_check"].items():
        col = column_index_from_string(letter)
        found = block.headings.get(col, "")
        if found not in accepted:
            warns.append(f"heading '{accepted[0]}' is not in column {letter} (found '{found}'); "
                         "the RE, PCR and HiFi sheets read fixed columns, check them.")
    return warns


def safe_name(project: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]", "_", project)


def write_protocol(wb, out: Path, rows: list[dtab.Row], project: str, cfg: dict) -> None:
    ws = wb[cfg["data_sheet"]]
    width = max(cfg["columns"].values())
    for r in range(2, ws.max_row + 1):
        for c in range(1, width + 1):
            ws.cell(row=r, column=c).value = None
    for i, row in enumerate(rows, start=2):
        for name, col in cfg["columns"].items():
            value = row.values.get(name)
            if name == "project_no":
                value = project
            elif name == "sub_no" and not row.own:
                value = None
            cell = ws.cell(row=i, column=col)
            try:
                cell.value = value
            except IllegalCharacterError:
                raise DesignError(f"line {row.line}, {name}: the value holds a control character "
                                  "that a workbook cannot store. Remove it.") from None
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"          # text typed as =... must not become a formula
    for sheet in wb.worksheets:
        for line in sheet.iter_rows():
            for cell in line:
                if cell.comment is not None:
                    cell.comment.width, cell.comment.height = COMMENT_SIZE
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    wb.calculation.fullCalcOnLoad = True
    wb.active = wb.sheetnames.index(cfg["data_sheet"])
    try:
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out)
    except OSError as exc:
        raise DesignError(f"cannot write {out}: {exc.strerror or exc}. If the file is open "
                          "in Excel, close it.") from None


def run(args: argparse.Namespace, say=print) -> int:
    cfg = load_config(Path(args.config) if args.config else None)
    if args.method is not None and not args.method.strip():
        raise DesignError("--method is empty. Give HiFi or GGA, or leave the option out.")
    wanted_method = resolve_method(cfg, args.method) if args.method else None
    date = args.date or dt.date.today().strftime("%y%m%d")
    try:
        if not re.fullmatch(r"\d{6}", date):
            raise ValueError(date)
        dt.datetime.strptime(date, "%y%m%d")  # noqa: DTZ007
    except ValueError:
        raise DesignError(f"--date must be a real date as yymmdd: '{date}'") from None
    if args.project is not None and not args.project.strip():
        raise DesignError("--project is empty.")
    wanted_subs = dtab.split_subs(args.sub)
    warnings: list[str] = []

    design = Path(args.design)
    if not design.is_file():
        raise DesignError(f"design table not found (or not a file): {design}")
    block = None
    if design.suffix.lower() in (".xlsx", ".xlsm"):
        sheet = args.sheet or cfg["design_sheet"]
        blocks = dtab.read_blocks(design, sheet, cfg)
        block = dtab.pick_block(blocks, args.project, cfg["construction_label"], sheet)
        project = block.project or args.project.strip().upper()
        rows = block.rows
        say(f"BLOCK {project} header=row{block.header} data=rows{block.first}-{block.last}")
        warnings += header_warnings(block, cfg)
        warnings += [f"in the block at row {block.header}: {note}" for note in block.stray]
        if block.unsaved:
            shown = ", ".join(block.unsaved[:5]) + (" ..." if len(block.unsaved) > 5 else "")
            warnings.append(
                f"{len(block.unsaved)} formula cell(s) have no saved value and are written "
                f"empty ({shown}). Open the workbook in Excel and save it, then run again.")
        csv_columns = None
    else:
        rows, csv_columns = dtab.read_csv(design, cfg)
        if not rows:
            raise DesignError(f"{design} has no data rows")
        projects = []
        for row in rows:
            name = row.text("project_no")
            if name and name.upper() not in [p.upper() for p in projects]:
                projects.append(name)
        if args.project and projects:
            wanted = args.project.strip().upper()
            blank = [str(r.line) for r in rows if not r.text("project_no")]
            rows = [r for r in rows if r.text("project_no").upper() == wanted]
            if blank:
                say(f"NOTE line(s) {', '.join(blank)} have no project_no and are left out "
                    f"because --project {args.project} is given.")
            if not rows:
                raise DesignError(f"Project {args.project} not found in {design}. "
                                  f"Projects present: {', '.join(projects)}")
            project = rows[0].text("project_no")      # the spelling of the table
        elif args.project:
            project = args.project.strip()
        elif len(projects) == 1:
            project = projects[0]
        elif not projects:
            raise DesignError("the CSV has no project_no values; pass --project.")
        else:
            raise DesignError(f"the CSV has several Project Nos ({', '.join(projects)}); "
                              "pass --project.")
        dtab.assign_subs(rows)

    method_notes: list[str] = []
    if wanted_method is not None:          # first, so that the checks below see only these rows
        rows, other = keep_method(rows, wanted_method)
        if other:
            method_notes.append(f"line(s) {', '.join(str(r.line) for r in other)} are left out: "
                                f"their method is not {wanted_method}.")
        if not rows:
            raise DesignError(f"no row of the table is for {wanted_method}.")
    if not wanted_subs:
        warnings += structure_warnings(rows)
    picked, notes = dtab.select_rows(rows, wanted_subs)
    notes = method_notes + notes
    for row in picked:
        say(f"ROW {row.line}  sub={row.sub}  name={row.text('plasmid_name')}  "
            f"RE={row.text('re1')}/{row.text('re2')}  backbone={row.text('backbone')}  "
            f"primers={row.text('oligo1')}/{row.text('oligo2')}  template={row.text('template')}")
    for note in notes:
        say(f"NOTE {note}")

    seen = sorted({r.text("method").lower(): r.text("method") for r in picked
                   if r.text("method")}.values())
    if wanted_method is None:
        if not seen:
            method = cfg["default_method"]
            say(f"NOTE the method column is empty; using {method}. Pass --method to choose.")
        elif len(seen) > 1:
            raise DesignError(f"the table mixes methods ({', '.join(seen)}); "
                              "pass --method HiFi or --method GGA to keep only that method's rows.")
        else:
            method = resolve_method(cfg, seen[0])
    else:
        method = wanted_method
    spec = cfg["methods"][method]

    if csv_columns is not None:
        missing = [c for c in spec["required_columns"] if c not in csv_columns]
        if missing:
            guess = wanted_method is None and not seen      # the method was only assumed
            hint = " For a GGA table pass --method GGA." if guess else ""
            raise DesignError(f"the CSV lacks column(s) for {method}: {', '.join(missing)}.{hint}")
    dtab.coerce(picked, spec["numeric_columns"])

    template = Path(args.template) if args.template else ROOT / spec["template"]
    if not template.is_file():
        raise DesignError(f"template not found: {template}")
    wb = open_template(template, cfg, spec, method)
    out = Path(args.out_dir) / f"{date}.{safe_name(project)}.{spec['file_suffix']}.xlsx"
    say(f"METHOD {method}  template={template}  max rows={spec['max_rows']}")
    if not args.dry_run and out.exists() and not args.force:
        raise DesignError(f"output already exists: {out}  (--force overwrites, or choose "
                          "another --date)")
    if len(picked) > spec["max_rows"]:
        warnings.append(f"{len(picked)} data rows; the {method} template covers only "
                        f"{spec['max_rows']} ({spec['max_rows_note']}). Rows beyond that are "
                        "not in the protocol.")
    if method == "HiFi":
        found = read_relist(wb)
        if found is None:
            say("NOTE the template has no RElist sheet; enzymes are not checked.")
        else:
            table, relist_warnings = found
            warnings += relist_warnings + check_relist(picked, table)
    for message in warnings:
        say(f"WARN {message}")
    if args.dry_run:
        if out.exists() and not args.force:
            say(f"NOTE {out} exists; a real run stops there unless --force is given.")
        say(f"DRYRUN would write {out}")
        return 0
    write_protocol(wb, out, picked, project, cfg)
    say(f"SAVED {out}  ({len(picked)} data rows)")
    return 0


def build_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--design", required=True, help="CSV file, or project workbook (.xlsx)")
    ap.add_argument("--sheet", help="sheet of the project workbook (default: config design_sheet)")
    ap.add_argument("--project", help="Project No. (needed when the table has several)")
    ap.add_argument("--method",
                    help="HiFi or GGA; keeps only the rows of that method "
                         "(default: the method column, else HiFi)")
    ap.add_argument("--sub", action="append",
                    help="only these Sub. No., each with its continuation rows (3,4,5)")
    ap.add_argument("--date", help="file date prefix, yymmdd (default: today)")
    ap.add_argument("--template", help="template workbook (default: the bundled one)")
    ap.add_argument("--out-dir", default="output", help="output folder (default: ./output)")
    ap.add_argument("--config", help="config.json (default: the one beside scripts/)")
    ap.add_argument("--dry-run", action="store_true", help="show the plan, write nothing")
    ap.add_argument("--force", action="store_true", help="overwrite an existing output file")
    return ap


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(errors="replace")
    try:
        return run(args)
    except DesignError as exc:
        print(f"ERROR {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
