"""make_protocol.py: the command line, the warnings and what is written."""

from __future__ import annotations

import re
import zipfile

import make_protocol as mp
import openpyxl
from conftest import HEADINGS, NAMES, hifi_row, write_csv, write_project_sheet

DATE = "261011"


def run(capsys, *args):
    code = mp.main([str(a) for a in args])
    out = capsys.readouterr()
    return code, out.out, out.err


def data_sheet(path):
    return openpyxl.load_workbook(path)["Data"]


def lines(out, tag):
    return [ln for ln in out.splitlines() if ln.startswith(tag)]


def test_xlsx_block_to_protocol(capsys, tmp_path, project_xlsx, enzyme_template):
    code, out, _ = run(capsys, "--design", project_xlsx, "--project", "P001", "--date", DATE,
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0
    target = tmp_path / "o" / f"{DATE}.P001.HiFi.TF.xlsx"
    assert target.exists() and "SAVED" in out and "(4 data rows)" in out
    assert lines(out, "WARN") == []
    ws = data_sheet(target)
    assert ws["A1"].value == "Plasmid Construction"          # the template heading stays
    assert [ws.cell(row=r, column=3).value for r in range(2, 7)] == [1, 2, 3, None, None]
    assert ws["H2"].value == "EcoRI-HF" and ws["I2"].value == "BamHI-HF"
    assert ws["B5"].value == "P001"                           # Project No. is on every row
    assert ws["M5"].value == "O-7" and ws["P5"].value == 500
    assert ws["G5"].value is None


def test_csv_gives_the_same_data_as_xlsx(capsys, tmp_path, project_xlsx, design_csv,
                                         enzyme_template):
    for src, name in ((project_xlsx, "x"), (design_csv, "c")):
        code, *_ = run(capsys, "--design", src, "--project", "P001", "--date", DATE,
                       "--template", enzyme_template, "--out-dir", tmp_path / name)
        assert code == 0
    a = data_sheet(tmp_path / "x" / f"{DATE}.P001.HiFi.TF.xlsx")
    b = data_sheet(tmp_path / "c" / f"{DATE}.P001.HiFi.TF.xlsx")
    for col in (2, 3, 7, 8, 9, 10, 12, 13, 14, 15, 16):
        for r in range(2, 8):
            va, vb = a.cell(row=r, column=col).value, b.cell(row=r, column=col).value
            assert (type(va), va) == (type(vb), vb), (r, col)


def test_example_csv_runs_without_warnings_once_enzymes_are_in_relist(
        capsys, tmp_path, root, enzyme_template):
    code, out, _ = run(capsys, "--design", root / "examples" / "design_table.csv",
                       "--date", DATE, "--template", enzyme_template, "--out-dir", tmp_path)
    assert code == 0 and lines(out, "WARN") == []
    assert (tmp_path / f"{DATE}.EX001.HiFi.TF.xlsx").exists()


def test_bundled_relist_is_empty_so_every_enzyme_warns(capsys, tmp_path, root):
    code, out, _ = run(capsys, "--design", root / "examples" / "design_table.csv",
                       "--date", DATE, "--out-dir", tmp_path)
    assert code == 0
    assert any("re1 'EcoRI-HF' is not in the RElist" in w for w in lines(out, "WARN"))


def test_enzyme_warnings(capsys, tmp_path, rows, enzyme_template):
    rows[0].update(re1="O-99", re2="")                    # primer-like re1: a PCR backbone
    rows[2].update(re1="BsmI", re2="NoSuchI")             # known re1, unknown re2
    rows[1].update(re1="BglII")                           # in the table, but not CutSmart
    design = write_csv(tmp_path / "d.csv", rows)
    code, out, _ = run(capsys, "--design", design, "--date", DATE,
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0
    warns = " ".join(lines(out, "WARN"))
    assert "re1 'O-99' is not in the RElist" in warns and "backbone prepared by PCR" in warns
    assert "re2 'NoSuchI' is not in the RElist" in warns and "'Check'" in warns
    assert "built for 10ﾃ佑utSmart" in warns and "BglII" in warns and "r3.1" in warns


def test_a_pcr_backbone_row_gets_one_warning_not_two(capsys, tmp_path, rows, enzyme_template):
    rows[0].update(re1="O-100", re2="O-101")             # primer names: the backbone is a PCR
    design = write_csv(tmp_path / "d.csv", rows)
    _, out, _ = run(capsys, "--design", design, "--date", DATE, "--template", enzyme_template,
                    "--out-dir", tmp_path / "o")
    warns = lines(out, "WARN")
    assert len(warns) == 1 and "re1 'O-100'" in warns[0] and "re2" not in warns[0]


def test_enzyme_names_are_matched_ignoring_case(capsys, tmp_path, rows, enzyme_template):
    rows[0].update(re1="ecori-hf", re2="BAMHI-HF")
    design = write_csv(tmp_path / "d.csv", rows)
    _, out, _ = run(capsys, "--design", design, "--date", DATE, "--template", enzyme_template,
                    "--out-dir", tmp_path / "o")
    assert lines(out, "WARN") == []


def test_too_many_rows_warn(capsys, tmp_path, enzyme_template):
    many = [hifi_row(i, f"p{i}", "EcoRI-HF", "", "bb", 100, "O-1", "O-2", "t", 10)
            for i in range(1, 14)]
    code, out, _ = run(capsys, "--design", write_csv(tmp_path / "d.csv", many), "--date", DATE,
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0
    assert any("13 data rows" in w and "only 12" in w for w in lines(out, "WARN"))


def test_dry_run_writes_nothing(capsys, tmp_path, design_csv, enzyme_template):
    code, out, _ = run(capsys, "--design", design_csv, "--date", DATE, "--dry-run",
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0 and "DRYRUN would write" in out
    assert not (tmp_path / "o").exists()


def test_existing_output_is_kept_unless_forced(capsys, tmp_path, design_csv, enzyme_template):
    args = ["--design", design_csv, "--date", DATE, "--template", enzyme_template,
            "--out-dir", tmp_path / "o"]
    assert run(capsys, *args)[0] == 0
    code, _, err = run(capsys, *args)
    assert code == 1 and "already exists" in err
    assert run(capsys, *args, "--force")[0] == 0


def test_default_output_folder_is_output_under_the_working_directory(
        capsys, tmp_path, monkeypatch, design_csv, enzyme_template):
    monkeypatch.chdir(tmp_path)
    code, *_ = run(capsys, "--design", design_csv, "--date", DATE, "--template", enzyme_template)
    assert code == 0 and (tmp_path / "output" / f"{DATE}.P001.HiFi.TF.xlsx").exists()


def test_project_name_is_made_safe_for_a_file_name(capsys, tmp_path, rows, enzyme_template):
    for r in rows:
        r["project_no"] = "Proj A/1"
    code, *_ = run(capsys, "--design", write_csv(tmp_path / "d.csv", rows), "--date", DATE,
                   "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0 and (tmp_path / "o" / f"{DATE}.Proj_A_1.HiFi.TF.xlsx").exists()


def test_bad_date(capsys, design_csv):
    code, _, err = run(capsys, "--design", design_csv, "--date", "2026-10-11")
    assert code == 1 and "yymmdd" in err


def test_missing_design_file(capsys, tmp_path):
    code, _, err = run(capsys, "--design", tmp_path / "none.csv")
    assert code == 1 and "not found" in err


def test_mixed_methods_need_the_option(capsys, tmp_path, rows):
    rows[0]["method"] = "GGA"
    code, _, err = run(capsys, "--design", write_csv(tmp_path / "d.csv", rows))
    assert code == 1 and "mixes methods" in err and "GGA, HiFi" in err


def test_gga_method_from_the_table(capsys, tmp_path):
    gga = [hifi_row(i, f"g{i}", "BsaI", "", "vec", None, f"O-{i}a", f"O-{i}b", "", None,
                    project="G1", method="GGA") for i in (1, 2, 3)]
    code, out, _ = run(capsys, "--design", write_csv(tmp_path / "g.csv", gga), "--date", DATE,
                       "--out-dir", tmp_path / "o")
    assert code == 0 and "METHOD GGA" in out and lines(out, "WARN") == []
    ws = data_sheet(tmp_path / "o" / f"{DATE}.G1.GGA.TF.xlsx")
    assert ws["H3"].value == "BsaI" and ws["M4"].value == "O-3a"


def test_the_method_option_overrides_and_warns_on_conflict(capsys, tmp_path, rows,
                                                           enzyme_template):
    design = write_csv(tmp_path / "d.csv", rows)
    code, out, _ = run(capsys, "--design", design, "--method", "hifi", "--date", DATE,
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0 and "METHOD HiFi" in out
    assert run(capsys, "--design", design, "--method", "PCR")[0] == 1


def test_empty_method_column_uses_hifi_and_says_so(capsys, tmp_path, rows, enzyme_template):
    for r in rows:
        r["method"] = None
    code, out, _ = run(capsys, "--design", write_csv(tmp_path / "d.csv", rows), "--date", DATE,
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0 and "using HiFi" in out


def test_sub_option_takes_continuation_rows(capsys, tmp_path, design_csv, enzyme_template):
    code, out, _ = run(capsys, "--design", design_csv, "--sub", "3", "--date", DATE,
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0 and "(2 data rows)" in out
    ws = data_sheet(tmp_path / "o" / f"{DATE}.P001.HiFi.TF.xlsx")
    assert ws["C2"].value == 3 and ws["C3"].value is None and ws["M3"].value == "O-7"
    assert ws["B3"].value == "P001" and ws["H4"].value is None


def test_missing_csv_columns_are_named(capsys, tmp_path, rows):
    header = [n for n in NAMES if n not in ("template", "insert_length")]
    code, _, err = run(capsys, "--design", write_csv(tmp_path / "d.csv", rows, header=header))
    assert code == 1 and "template, insert_length" in err


def test_csv_number_that_is_not_a_number(capsys, tmp_path, rows):
    rows[1]["backbone_length"] = "5.2 kb"
    code, _, err = run(capsys, "--design", write_csv(tmp_path / "d.csv", rows))
    assert code == 1 and "line 3" in err and "backbone_length" in err


def test_project_option_picks_among_several_projects_in_a_csv(capsys, tmp_path, rows,
                                                              enzyme_template):
    other = [hifi_row(1, "q", "EcoRI-HF", "", "bb", 10, "O-1", "O-2", "t", 10, project="P002")]
    design = write_csv(tmp_path / "d.csv", rows + other)
    assert run(capsys, "--design", design)[0] == 1
    code, out, _ = run(capsys, "--design", design, "--project", "p002", "--date", DATE,
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0 and "(1 data rows)" in out
    code, _, err = run(capsys, "--design", design, "--project", "P009")
    assert code == 1 and "P001, P002" in err


def test_csv_without_a_project_column_needs_the_option(capsys, tmp_path, rows, enzyme_template):
    header = [n for n in NAMES if n != "project_no"]
    design = write_csv(tmp_path / "d.csv", rows, header=header)
    code, _, err = run(capsys, "--design", design)
    assert code == 1 and "--project" in err
    code, *_ = run(capsys, "--design", design, "--project", "P7", "--date", DATE,
                   "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0 and (tmp_path / "o" / f"{DATE}.P7.HiFi.TF.xlsx").exists()


def test_xlsx_without_a_project_value_uses_the_option(capsys, tmp_path, rows, enzyme_template):
    for r in rows:
        r["project_no"] = None
    path = write_project_sheet(tmp_path / "p.xlsx", [rows])
    assert run(capsys, "--design", path)[0] == 1
    code, *_ = run(capsys, "--design", path, "--project", "P555", "--date", DATE,
                   "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0
    assert data_sheet(tmp_path / "o" / f"{DATE}.P555.HiFi.TF.xlsx")["B2"].value == "P555"


def test_sheet_option(capsys, tmp_path, rows, enzyme_template):
    path = write_project_sheet(tmp_path / "p.xlsx", [rows], sheet="Records")
    assert run(capsys, "--design", path)[0] == 1          # default sheet name is Projects
    code, *_ = run(capsys, "--design", path, "--sheet", "Records", "--date", DATE,
                   "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0


def test_heading_warnings_accept_both_spellings(capsys, tmp_path, rows, enzyme_template):
    for spelling in ("Backbone plasmid/gDNA", "Backborn plasmid/gDNA"):
        headings = list(HEADINGS)
        headings[9] = spelling
        path = write_project_sheet(tmp_path / "p.xlsx", [rows], headings=headings)
        _, out, _ = run(capsys, "--design", path, "--date", DATE, "--force",
                        "--template", enzyme_template, "--out-dir", tmp_path / "o")
        assert lines(out, "WARN") == []
    headings = list(HEADINGS)
    headings[12] = "Primer A"
    path = write_project_sheet(tmp_path / "q.xlsx", [rows], headings=headings)
    _, out, _ = run(capsys, "--design", path, "--date", DATE, "--force",
                    "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert any("not in column M" in w for w in lines(out, "WARN"))


def test_old_data_in_the_template_is_cleared(capsys, tmp_path, design_csv, enzyme_template):
    wb = openpyxl.load_workbook(enzyme_template)
    wb["Data"]["H9"] = "left over"
    wb["Data"]["AG12"] = "left over"
    wb.save(enzyme_template)
    run(capsys, "--design", design_csv, "--date", DATE, "--template", enzyme_template,
        "--out-dir", tmp_path / "o")
    ws = data_sheet(tmp_path / "o" / f"{DATE}.P001.HiFi.TF.xlsx")
    assert ws["H9"].value is None and ws["AG12"].value is None


def test_template_without_relist_is_not_checked(capsys, tmp_path, design_csv, enzyme_template):
    wb = openpyxl.load_workbook(enzyme_template)
    del wb["RElist"]
    wb.save(enzyme_template)
    code, out, _ = run(capsys, "--design", design_csv, "--date", DATE,
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0 and "no RElist sheet" in out


def test_saved_workbook_carries_no_author_and_full_comment_size(
        capsys, tmp_path, design_csv, enzyme_template):
    run(capsys, "--design", design_csv, "--date", DATE, "--template", enzyme_template,
        "--out-dir", tmp_path / "o")
    target = tmp_path / "o" / f"{DATE}.P001.HiFi.TF.xlsx"
    with zipfile.ZipFile(target) as z:
        core = z.read("docProps/core.xml").decode("utf-8")
        vml = "".join(z.read(n).decode("utf-8") for n in z.namelist()
                      if n.startswith("xl/drawings/") and n.endswith(".vml"))
        book = z.read("xl/workbook.xml").decode("utf-8")
    # an empty element, written either as <x></x> or as <x/>, whichever XML writer is installed
    assert re.search(r"<dc:creator[^>]*(/>|></dc:creator>)", core)
    assert re.search(r"<cp:lastModifiedBy\s*(/>|></cp:lastModifiedBy>)", core)
    assert "width:300px;height:120px" in vml
    assert 'fullCalcOnLoad="1"' in book
    props = openpyxl.load_workbook(target).properties
    assert props.creator in (None, "") and props.lastModifiedBy in (None, "")
