"""Offline fixtures: design tables built in-process with openpyxl.

No Excel, no network. Workbooks are generated at run time because *.xlsx files
are not tracked (only the two bundled templates are).
"""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

CONFIG = json.loads((ROOT / "config.json").read_text(encoding="utf-8"))
NAMES = sorted(CONFIG["columns"], key=CONFIG["columns"].get)

HEADINGS = [
    "Plasmid Construction", "Project No.", "Sub. No.", "Method", "Plasmid map", "Length",
    "Plasmid name", "RE1", "RE2", "Backbone plasmid/gDNA", "Backbone prep.", "Length (bp)",
    "PCR primer/Oligo 1", "PCR primer/Oligo 2", "Template", "Length (bp)", "PCR/Annealing",
    "DpnI", "Transformation", "Miniprep", "Plasmid Ext.", "Nanopore", "Sequence check",
    "Plasmid No.", "Plasmid name", "Competent cell", "Conc", "Design", "Order", "Obtained",
    "Made by", "Entered by", "Project number",
]

# Enzymes a test adds to RElist: name, buffer, reaction temperature, heat inactivation.
ENZYMES = [
    ("EcoRI-HF", "CutSmart", "37°C", "65°C for 20 minutes"),
    ("BamHI-HF", "CutSmart", "37°C", "No"),
    ("XhoI", "CutSmart", "37°C", "65°C for 20 minutes"),
    ("BsmI", "CutSmart", "65°C", "80°C for 20 minutes"),
    ("BglII", "r3.1", "37°C", "No"),
]


def hifi_row(sub, name="", re1="", re2="", backbone="", bb_len=None, o1="", o2="",
             template="", ins_len=None, project="P001", method="HiFi"):
    """One design row as a dict of config column names."""
    row = dict.fromkeys(NAMES)
    row.update(project_no=project, sub_no=sub if sub != "" else None, method=method or None,
               plasmid_name=name or None,
               re1=re1 or None, re2=re2 or None, backbone=backbone or None,
               backbone_length=bb_len, oligo1=o1 or None, oligo2=o2 or None,
               template=template or None, insert_length=ins_len)
    return row


DEFAULT_ROWS = [
    hifi_row(1, "p-one", "EcoRI-HF", "BamHI-HF", "bb1", 5000, "O-1", "O-2", "t1", 1000),
    hifi_row(2, "p-two", "XhoI", "", "bb2", 5200, "O-3", "O-4", "t2", 800),
    hifi_row(3, "p-three", "BsmI", "EcoRI-HF", "bb3", 4900, "O-5", "O-6", "t3", 700),
    hifi_row("", "", "", "", "", None, "O-7", "O-8", "t4", 500),
]


def write_project_sheet(path, blocks, sheet="Projects", headings=None,
                        label="Plasmid Construction"):
    """A project workbook: each block is a list of rows under a header row, a blank row between."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    r = 1
    for rows in blocks:
        for c, h in enumerate(headings or HEADINGS, 1):
            ws.cell(row=r, column=c, value=h)
        ws.cell(row=r, column=1, value=label)
        for row in rows:
            r += 1
            for name in NAMES:
                value = row.get(name)
                if value is not None:
                    ws.cell(row=r, column=CONFIG["columns"][name], value=value)
        r += 2
    wb.save(path)
    return path


def write_csv(path, rows, header=None):
    header = header or NAMES
    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(header)
        for row in rows:
            w.writerow(["" if row.get(n) is None else row.get(n) for n in header])
    return path


@pytest.fixture
def root():
    return ROOT


@pytest.fixture
def rows():
    return [dict(r) for r in DEFAULT_ROWS]


@pytest.fixture
def project_xlsx(tmp_path, rows):
    return write_project_sheet(tmp_path / "projects.xlsx", [rows])


@pytest.fixture
def design_csv(tmp_path, rows):
    return write_csv(tmp_path / "design.csv", rows)


@pytest.fixture
def enzyme_template(tmp_path):
    """The bundled HiFi template with a few enzymes added to RElist."""
    wb = openpyxl.load_workbook(ROOT / "templates" / "protocol.hifi.xlsx")
    ws = wb["RElist"]
    for i, row in enumerate(ENZYMES, start=4):
        for c, value in enumerate(row, start=1):
            ws.cell(row=i, column=c, value=value)
    path = tmp_path / "template.hifi.xlsx"
    wb.save(path)
    return path
