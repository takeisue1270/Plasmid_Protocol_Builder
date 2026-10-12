"""config.json: the column map, the two methods and the template notes."""

from __future__ import annotations

import openpyxl
import pytest
from conftest import CONFIG, ROOT
from openpyxl.utils import column_index_from_string

COLUMNS = CONFIG["columns"]


def test_the_column_map_covers_b_to_ag_once():
    assert sorted(COLUMNS.values()) == list(range(2, 34))
    assert len(COLUMNS) == 32


@pytest.mark.parametrize("method", ["HiFi", "GGA"])
def test_method_settings_use_known_columns(method):
    spec = CONFIG["methods"][method]
    for key in ("read_columns", "required_columns", "numeric_columns"):
        assert set(spec[key]) <= set(COLUMNS), (method, key)
    assert set(spec["numeric_columns"]) <= set(spec["read_columns"])
    assert "sub_no" in spec["required_columns"]
    assert (ROOT / spec["template"]).exists()
    assert spec["required_sheet"] in openpyxl.load_workbook(ROOT / spec["template"]).sheetnames
    assert spec["max_rows"] > 0 and spec["file_suffix"].endswith(".TF")


def test_read_columns_of_gga_are_part_of_those_of_hifi():
    assert set(CONFIG["methods"]["GGA"]["read_columns"]) <= set(
        CONFIG["methods"]["HiFi"]["read_columns"])


def test_default_method_is_one_of_the_methods():
    assert CONFIG["default_method"] in CONFIG["methods"]


def test_heading_check_points_at_the_columns_make_protocol_reads():
    expected = {"J": "backbone", "M": "oligo1", "N": "oligo2", "O": "template"}
    for letter, name in expected.items():
        assert COLUMNS[name] == column_index_from_string(letter)
    assert CONFIG["header_check"]["J"][0] == "Backbone plasmid/gDNA"
    assert all(isinstance(v, list) and v for v in CONFIG["header_check"].values())


def test_template_notes_are_well_formed_and_point_at_cells_that_exist():
    books = {"HiFi": openpyxl.load_workbook(ROOT / "templates" / "protocol.hifi.xlsx"),
             "GGA": openpyxl.load_workbook(ROOT / "templates" / "protocol.gga.xlsx")}
    seen = set()
    for note in CONFIG["template_notes"]:
        assert set(note) == {"template", "sheet", "cell", "text"}
        assert note["template"] in ("HiFi", "GGA", "both")
        assert note["text"].strip()
        for kind in (("HiFi", "GGA") if note["template"] == "both" else (note["template"],)):
            assert note["sheet"] in books[kind].sheetnames, note
            assert books[kind][note["sheet"]][note["cell"]].value is not None, note
            assert (kind, note["sheet"], note["cell"]) not in seen
            seen.add((kind, note["sheet"], note["cell"]))
