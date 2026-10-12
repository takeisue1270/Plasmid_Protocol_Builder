"""Bad input: each case ends in a message that names the cause, never a traceback."""

from __future__ import annotations

import openpyxl
from conftest import HEADINGS, hifi_row, write_csv, write_project_sheet
from test_make_protocol import DATE, lines, run


def test_row_before_the_first_sub_is_warned(capsys, tmp_path, rows, enzyme_template):
    rows[0]["sub_no"] = None
    design = write_csv(tmp_path / "d.csv", rows)
    code, out, _ = run(capsys, "--design", design, "--date", DATE, "--template", enzyme_template,
                       "--out-dir", tmp_path / "o")
    assert code == 0
    assert any("line(s) 2 have no Sub. No." in w for w in lines(out, "WARN"))


def test_a_sub_used_twice_is_warned_but_a_continuation_row_is_not(capsys, tmp_path, rows,
                                                                  enzyme_template):
    _, out, _ = run(capsys, "--design", write_csv(tmp_path / "a.csv", rows), "--date", DATE,
                    "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert lines(out, "WARN") == []
    rows[1]["sub_no"] = 1
    _, out, _ = run(capsys, "--design", write_csv(tmp_path / "b.csv", rows), "--date", DATE,
                    "--template", enzyme_template, "--out-dir", tmp_path / "o2")
    assert any("Sub. No. 1 is used on lines 2, 3" in w for w in lines(out, "WARN"))


def test_workbook_numbers_are_checked_like_csv_numbers(capsys, tmp_path, rows, enzyme_template):
    rows[1]["insert_length"] = "1.2 kb"
    path = write_project_sheet(tmp_path / "p.xlsx", [rows])
    code, _, err = run(capsys, "--design", path, "--date", DATE, "--template", enzyme_template,
                       "--out-dir", tmp_path / "o")
    assert code == 1 and "insert_length" in err and "1.2 kb" in err
    rows[1]["insert_length"] = "1200"                      # text that reads as a number
    path = write_project_sheet(tmp_path / "q.xlsx", [rows])
    code, *_ = run(capsys, "--design", path, "--date", DATE, "--template", enzyme_template,
                   "--out-dir", tmp_path / "o")
    assert code == 0
    ws = openpyxl.load_workbook(tmp_path / "o" / f"{DATE}.P001.HiFi.TF.xlsx")["Data"]
    assert ws["P3"].value == 1200


def test_a_csv_not_in_utf8_is_explained(capsys, tmp_path, rows):
    path = write_csv(tmp_path / "d.csv", rows)
    path.write_bytes(path.read_bytes().replace(b"p-one", bytes([0x83, 0x76, 0x83, 0x89])))
    code, _, err = run(capsys, "--design", path)
    assert code == 1 and "UTF-8" in err and "Traceback" not in err


def test_a_csv_with_other_headings_names_the_accepted_ones(capsys, tmp_path):
    path = tmp_path / "d.csv"
    path.write_text("Project No.;Sub. No.;RE1\nP1;1;EcoRI\n", encoding="utf-8")
    code, _, err = run(capsys, "--design", path)
    assert code == 1 and "sub_no" in err and "plasmid_name" in err
    path.write_text("", encoding="utf-8")
    assert run(capsys, "--design", path)[0] == 1


def test_sub_written_as_a_decimal_is_still_found(capsys, tmp_path, rows, enzyme_template):
    rows[0]["sub_no"] = "1.0"
    design = write_csv(tmp_path / "d.csv", rows)
    code, out, _ = run(capsys, "--design", design, "--sub", "1", "--date", DATE,
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0 and "(1 data rows)" in out


def test_project_keeps_the_spelling_of_the_table(capsys, tmp_path, rows, enzyme_template):
    other = [hifi_row(1, "q", "EcoRI-HF", "", "bb", 10, "O-1", "O-2", "t", 10, project="P2")]
    design = write_csv(tmp_path / "d.csv", rows + other)
    code, *_ = run(capsys, "--design", design, "--project", "p2", "--date", DATE,
                   "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0 and (tmp_path / "o" / f"{DATE}.P2.HiFi.TF.xlsx").exists()
    assert openpyxl.load_workbook(tmp_path / "o" / f"{DATE}.P2.HiFi.TF.xlsx")["Data"]["B2"].value \
        == "P2"


def test_impossible_date(capsys, design_csv):
    code, _, err = run(capsys, "--design", design_csv, "--date", "261399")
    assert code == 1 and "real date" in err


def test_a_folder_is_not_a_design_file(capsys, tmp_path):
    code, _, err = run(capsys, "--design", tmp_path)
    assert code == 1 and "not a file" in err


def test_a_template_that_is_not_a_workbook(capsys, tmp_path, design_csv):
    bad = tmp_path / "t.xlsx"
    bad.write_text("not a zip", encoding="utf-8")
    code, _, err = run(capsys, "--design", design_csv, "--template", bad)
    assert code == 1 and "cannot read the template" in err


def test_a_template_of_the_other_method(capsys, tmp_path, root, design_csv):
    code, _, err = run(capsys, "--design", design_csv, "--template",
                       root / "templates" / "protocol.gga.xlsx")
    assert code == 1 and "no 'HiFi' sheet" in err and "Anneal" in err


def test_a_template_without_a_data_sheet(capsys, tmp_path, design_csv, enzyme_template):
    wb = openpyxl.load_workbook(enzyme_template)
    del wb["Data"]
    wb.save(enzyme_template)
    code, _, err = run(capsys, "--design", design_csv, "--template", enzyme_template)
    assert code == 1 and "no 'Data' sheet" in err


def test_an_output_that_cannot_be_written(capsys, tmp_path, design_csv, enzyme_template):
    blocker = tmp_path / "o" / f"{DATE}.P001.HiFi.TF.xlsx"
    blocker.mkdir(parents=True)                           # a folder where the file should go
    code, _, err = run(capsys, "--design", design_csv, "--date", DATE, "--force",
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 1 and "cannot write" in err


def test_a_row_without_re1_is_not_checked(capsys, tmp_path, rows, enzyme_template):
    rows[0].update(re1="", re2="NoSuchI")                  # the RE sheet leaves it blank
    _, out, _ = run(capsys, "--design", write_csv(tmp_path / "d.csv", rows), "--date", DATE,
                    "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert lines(out, "WARN") == []


def test_workbook_text_is_trimmed_like_csv_text(capsys, tmp_path, rows, enzyme_template):
    rows[0]["re1"] = "EcoRI-HF "                           # a trailing space typed in Excel
    path = write_project_sheet(tmp_path / "p.xlsx", [rows])
    _, out, _ = run(capsys, "--design", path, "--date", DATE, "--template", enzyme_template,
                    "--out-dir", tmp_path / "o")
    assert lines(out, "WARN") == []
    ws = openpyxl.load_workbook(tmp_path / "o" / f"{DATE}.P001.HiFi.TF.xlsx")["Data"]
    assert ws["H2"].value == "EcoRI-HF"


def test_a_design_workbook_that_is_not_a_workbook(capsys, tmp_path):
    bad = tmp_path / "p.xlsx"
    bad.write_text("not a zip", encoding="utf-8")
    code, _, err = run(capsys, "--design", bad)
    assert code == 1 and "cannot read" in err and "Traceback" not in err


def test_dates_need_six_digits(capsys, design_csv):
    for bad in ("26101", "2610111"):
        code, _, err = run(capsys, "--design", design_csv, "--date", bad)
        assert code == 1 and "yymmdd" in err


def test_dry_run_says_when_the_real_run_would_stop(capsys, tmp_path, design_csv, enzyme_template):
    args = ["--design", design_csv, "--date", DATE, "--template", enzyme_template,
            "--out-dir", tmp_path / "o"]
    run(capsys, *args)
    _, out, _ = run(capsys, *args, "--dry-run")
    assert "exists; a real run stops" in out
    _, out, _ = run(capsys, *args, "--dry-run", "--force")
    assert "exists" not in out


def test_a_zip_that_is_not_a_workbook(capsys, tmp_path):
    import zipfile
    bad = tmp_path / "p.xlsx"
    with zipfile.ZipFile(bad, "w") as z:
        z.writestr("hello.txt", "x")
    code, _, err = run(capsys, "--design", bad)
    assert code == 1 and "cannot read" in err and "Traceback" not in err


def test_relist_rows_that_the_re_sheet_cannot_use_are_warned(capsys, tmp_path, design_csv,
                                                             enzyme_template):
    wb = openpyxl.load_workbook(enzyme_template)
    ws = wb["RElist"]
    ws["A20"], ws["B20"], ws["C20"], ws["D20"] = "SpaceI ", "CutSmart", "37\u00b0C", "No"
    ws["A21"], ws["B21"] = "HalfI", "CutSmart"
    wb.save(enzyme_template)
    _, out, _ = run(capsys, "--design", design_csv, "--date", DATE, "--template", enzyme_template,
                    "--out-dir", tmp_path / "o")
    warns = " ".join(lines(out, "WARN"))
    assert "RElist row 20" in warns and "space at an end" in warns
    assert "RElist row 21 (HalfI): column(s) C, D are empty" in warns


def test_a_gga_table_without_a_method_column_says_how_to_run(capsys, tmp_path):
    header = ["sub_no", "plasmid_name", "re1", "backbone", "oligo1", "oligo2"]
    path = tmp_path / "g.csv"
    path.write_text(",".join(header) + "\n1,g,BsaI,vec,O-1,O-2\n", encoding="utf-8")
    code, _, err = run(capsys, "--design", path, "--project", "P1")
    assert code == 1 and "--method GGA" in err
    code, *_ = run(capsys, "--design", path, "--project", "P1", "--method", "GGA", "--date", DATE,
                   "--out-dir", tmp_path / "o")
    assert code == 0


def test_text_that_starts_with_an_equals_sign_stays_text(capsys, tmp_path, rows, enzyme_template):
    rows[0]["plasmid_name"] = "=1+1"
    code, *_ = run(capsys, "--design", write_csv(tmp_path / "d.csv", rows), "--date", DATE,
                   "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0
    ws = openpyxl.load_workbook(tmp_path / "o" / f"{DATE}.P001.HiFi.TF.xlsx")["Data"]
    assert ws["G2"].data_type == "s" and ws["G2"].value == "=1+1"


def test_a_missing_sub_lists_the_subs_that_exist(capsys, design_csv):
    code, _, err = run(capsys, "--design", design_csv, "--sub", "3-4")
    assert code == 1 and "Sub. Nos. present: 1, 2, 3" in err and "--sub 3,4,5" in err


def test_a_workbook_with_a_wrong_stored_dimension(capsys, tmp_path, rows, enzyme_template):
    import zipfile
    path = write_project_sheet(tmp_path / "p.xlsx", [rows])
    bad = tmp_path / "bad.xlsx"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(bad, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item == "xl/worksheets/sheet1.xml":
                import re
                data = re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="A1"', data)
            zout.writestr(item, data)
    code, out, _ = run(capsys, "--design", bad, "--project", "P001", "--date", DATE,
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0 and "(4 data rows)" in out


def test_a_sheet_without_a_block_says_what_a_block_is(capsys, tmp_path, rows):
    headings = list(HEADINGS)
    headings[1] = "Project No"                             # lacks the full stop
    path = write_project_sheet(tmp_path / "p.xlsx", [rows], headings=headings)
    for extra in ([], ["--project", "P001"]):
        code, _, err = run(capsys, "--design", path, *extra)
        assert code == 1 and "no block found" in err and "'Project No.'" in err


def test_method_spelling_is_not_a_conflict(capsys, tmp_path, rows):
    for i, r in enumerate(rows):
        r["method"] = "GGA" if i % 2 == 0 else "gga"
    code, out, err = run(capsys, "--design", write_csv(tmp_path / "d.csv", rows), "--date", DATE,
                         "--out-dir", tmp_path / "o")
    assert code == 0 and "METHOD GGA" in out and "mixes methods" not in err


def test_a_control_character_is_explained(capsys, tmp_path, rows, enzyme_template):
    rows[0]["plasmid_name"] = "bad\x01name"
    code, _, err = run(capsys, "--design", write_csv(tmp_path / "d.csv", rows), "--date", DATE,
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 1 and "control character" in err and "Traceback" not in err


def test_a_missing_config_is_explained(capsys, tmp_path, design_csv):
    code, _, err = run(capsys, "--design", design_csv, "--config", tmp_path / "none.json")
    assert code == 1 and "cannot read the configuration" in err


def test_formula_cells_without_a_saved_value_are_warned(capsys, tmp_path, rows, enzyme_template):
    path = write_project_sheet(tmp_path / "p.xlsx", [rows])
    wb = openpyxl.load_workbook(path)
    wb["Projects"]["L2"] = "=5000*1"                       # openpyxl saves no value for it
    wb.save(path)
    _, out, _ = run(capsys, "--design", path, "--date", DATE, "--template", enzyme_template,
                    "--out-dir", tmp_path / "o")
    assert any("no saved value" in w and "line 2 backbone_length" in w for w in lines(out, "WARN"))


def test_a_workbook_saved_by_excel_is_not_told_to_save(capsys, tmp_path, rows, enzyme_template):
    """Formula cells with a saved result, even an empty one, are not 'unsaved'."""
    import re
    import zipfile
    path = write_project_sheet(tmp_path / "p.xlsx", [rows])
    wb = openpyxl.load_workbook(path)
    wb["Projects"]["L2"] = "=5000*1"
    wb["Projects"]["I3"] = '=IF(1=1,"","x")'
    wb.save(path)
    saved = tmp_path / "saved.xlsx"
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(saved, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item == "xl/worksheets/sheet1.xml":
                # give the first formula a result; leave the second with an empty one
                data = re.sub(rb"(<f>5000\*1</f>)<v\s*(?:/>|></v>)", rb"\g<1><v>5000</v>", data)
            zout.writestr(item, data)
    _, out, _ = run(capsys, "--design", saved, "--date", DATE, "--template", enzyme_template,
                    "--out-dir", tmp_path / "o")
    assert not any("no saved value" in w for w in lines(out, "WARN"))


def test_project_number_formulas_are_not_counted_as_unsaved(capsys, tmp_path, rows,
                                                            enzyme_template):
    path = write_project_sheet(tmp_path / "p.xlsx", [rows])
    wb = openpyxl.load_workbook(path)
    wb["Projects"]["B3"] = "=B2"                           # Designer repeats the Project No.
    wb.save(path)
    _, out, _ = run(capsys, "--design", path, "--date", DATE, "--template", enzyme_template,
                    "--out-dir", tmp_path / "o")
    assert not any("no saved value" in w for w in lines(out, "WARN"))


def test_an_empty_project_option_and_a_foreign_config(capsys, tmp_path, design_csv):
    code, _, err = run(capsys, "--design", design_csv, "--project", " ")
    assert code == 1 and "--project is empty" in err
    odd = tmp_path / "odd.json"
    odd.write_text("[1]", encoding="utf-8")
    code, _, err = run(capsys, "--design", design_csv, "--config", odd)
    assert code == 1 and "not a configuration of this tool" in err


def test_rows_after_a_blank_row_are_reported(capsys, tmp_path, rows, enzyme_template):
    path = write_project_sheet(tmp_path / "p.xlsx", [rows])
    wb = openpyxl.load_workbook(path)
    ws = wb["Projects"]
    ws.delete_rows(4)
    ws.insert_rows(4)                                      # a blank row inside the block
    ws["B6"] = "P001"
    ws["C6"] = 4
    ws["G6"] = "late"
    wb.save(path)
    _, out, _ = run(capsys, "--design", path, "--date", DATE, "--template", enzyme_template,
                    "--out-dir", tmp_path / "o")
    assert any("follow a blank row" in w and "6" in w for w in lines(out, "WARN"))


def test_a_changing_project_number_inside_a_block_is_reported(capsys, tmp_path, rows,
                                                              enzyme_template):
    path = write_project_sheet(tmp_path / "p.xlsx", [rows])
    wb = openpyxl.load_workbook(path)
    wb["Projects"]["B4"] = "P099"
    wb.save(path)
    _, out, _ = run(capsys, "--design", path, "--date", DATE, "--template", enzyme_template,
                    "--out-dir", tmp_path / "o")
    assert any("changes inside the block (P099)" in w for w in lines(out, "WARN"))


def test_empty_sub_and_method_options_are_errors(capsys, design_csv):
    for option in ("--sub", "--method"):
        for empty in ("", ","):
            if option == "--method" and empty == ",":
                continue
            code, _, err = run(capsys, "--design", design_csv, option, empty)
            assert code == 1 and "is empty" in err, (option, empty)


def test_the_gga_hint_only_when_the_table_has_no_method(capsys, tmp_path):
    header = ["sub_no", "plasmid_name", "method", "re1"]
    path = tmp_path / "g.csv"
    path.write_text(",".join(header) + "\n1,g,GGA,BsaI\n", encoding="utf-8")
    code, _, err = run(capsys, "--design", path, "--project", "P1")
    assert code == 1 and "lacks column" in err and "--method GGA" not in err


def test_method_option_keeps_only_that_methods_rows(capsys, tmp_path, rows, enzyme_template):
    rows[1]["method"] = "GGA"                              # one construct of another method
    rows[3]["method"] = None                               # the continuation row of sub 3
    design = write_csv(tmp_path / "d.csv", rows)
    code, out, _ = run(capsys, "--design", design, "--method", "HiFi", "--date", DATE,
                       "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert code == 0 and "(3 data rows)" in out
    assert "line(s) 3 are left out" in out and "ROW 3 " not in out
    code, _, err = run(capsys, "--design", design, "--method", "GGA", "--sub", "1")
    assert code == 1 and "Sub. No. 1 not found" in err and "present: 2" in err
    rows[1]["method"] = "HiFi"
    code, _, err = run(capsys, "--design", write_csv(tmp_path / "h.csv", rows), "--method", "GGA")
    assert code == 1 and "no row of the table is for GGA" in err


def test_continuation_rows_follow_their_sub_when_a_method_is_dropped(capsys, tmp_path, rows,
                                                                    enzyme_template):
    for r in rows:
        r["method"] = None
    rows[2]["method"] = "GGA"                              # sub 3 and its continuation row
    design = write_csv(tmp_path / "d.csv", rows)
    _, out, _ = run(capsys, "--design", design, "--method", "HiFi", "--date", DATE,
                    "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert "(2 data rows)" in out and "line(s) 4, 5 are left out" in out


def test_rows_without_a_project_are_reported_when_project_is_given(capsys, tmp_path, rows,
                                                                   enzyme_template):
    rows[0]["project_no"] = None
    design = write_csv(tmp_path / "d.csv", rows)
    _, out, _ = run(capsys, "--design", design, "--project", "P001", "--date", DATE,
                    "--template", enzyme_template, "--out-dir", tmp_path / "o")
    assert "line(s) 2 have no project_no and are left out" in out


def test_sub_numbers_repeated_in_another_method_are_not_a_problem(capsys, tmp_path, rows,
                                                                  enzyme_template):
    gga = [hifi_row(i, f"g{i}", "BsaI", "", "vec", None, f"G-{i}a", f"G-{i}b", "", None,
                    method="GGA") for i in (1, 2)]
    for r in rows:
        r["method"] = "HiFi"
    design = write_csv(tmp_path / "d.csv", rows[:2] + gga)
    _, out, _ = run(capsys, "--design", design, "--method", "GGA", "--date", DATE,
                    "--out-dir", tmp_path / "o")
    assert not any("is used on lines" in w for w in lines(out, "WARN"))
    code, out, _ = run(capsys, "--design", design, "--method", "HiFi", "--sub", "1",
                       "--date", DATE, "--template", enzyme_template, "--out-dir", tmp_path / "p")
    assert code == 0 and "(1 data rows)" in out
