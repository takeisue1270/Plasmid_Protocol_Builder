"""The two bundled templates: English only, what they read, RElist, notes, golden files."""

from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import pytest
from conftest import CONFIG, ENZYMES, ROOT

sys.path.insert(0, str(ROOT / "scripts"))
import dump_template  # noqa: E402

# Kana, kanji, full-width forms, CJK symbols and punctuation (U+3000-303F, which holds the
# ideographic space and the ditto mark) and the degree-Celsius sign. The degree sign
# U+00B0, the multiplication sign and the registered-trademark sign are not in the range.
CJK_RX = re.compile(r"[\u3000-\u30ff\u31f0-\u31ff\u3400-\u9fff\uff00-\uffef\u2103]")

TEMPLATES = {"HiFi": ROOT / "templates" / "protocol.hifi.xlsx",
             "GGA": ROOT / "templates" / "protocol.gga.xlsx"}
SHEETS = {"HiFi": ["RElist", "Data", "RE", "PCR", "HiFi", "TF"],
          "GGA": ["Data", "Anneal", "GGA", "TF"]}


@pytest.fixture(params=["HiFi", "GGA"])
def method(request):
    return request.param


@pytest.fixture
def wb(method):
    return openpyxl.load_workbook(TEMPLATES[method])


def test_sheets(wb, method):
    assert wb.sheetnames == SHEETS[method]
    assert wb.active.title == "Data"
    assert wb.calculation.fullCalcOnLoad is True


def test_no_japanese_anywhere(wb):
    hits = []
    for ws in wb.worksheets:
        if CJK_RX.search(ws.title):
            hits.append((ws.title, "sheet name"))
        for row in ws.iter_rows():
            for c in row:
                where = f"{ws.title}!{c.coordinate}"
                if isinstance(c.value, str) and CJK_RX.search(c.value):
                    hits.append((where, c.value))
                if CJK_RX.search(c.number_format or ""):
                    hits.append((where, c.number_format))
                if c.comment is not None and CJK_RX.search(c.comment.text):
                    hits.append((where, c.comment.text))
        for rng, rules in ws.conditional_formatting._cf_rules.items():
            for rule in rules:
                for f in rule.formula or []:
                    if CJK_RX.search(str(f)):
                        hits.append((f"{ws.title}!{rng.sqref}", str(f)))
    for name in wb.defined_names:
        if CJK_RX.search(name):
            hits.append(("defined name", name))
    assert hits == []


def test_the_cjk_pattern_catches_what_the_originals_used():
    for sample in ("\u3007", "\u3003", "a\u3000b", "\u2103", "\u6e29\u5ea6", "\u30c6", "\uff1a"):
        assert CJK_RX.search(sample), sample
    for sample in ("°C", "×", "®", "µL", "→", "□", "α"):
        assert not CJK_RX.search(sample), sample


def test_formulas_read_exactly_the_columns_the_config_lists(wb, method):
    letters = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    letters.update(re.findall(r"Data!\$?([A-Z]+)\$?\d+", c.value))
    expected = {openpyxl.utils.get_column_letter(CONFIG["columns"][n])
                for n in CONFIG["methods"][method]["read_columns"]}
    assert letters == expected


def test_sub_number_column_is_formatted_as_two_digits(wb, method):
    last = CONFIG["methods"][method]["max_rows"] + 1
    ws = wb["Data"]
    assert all(ws[f"C{r}"].number_format == "00" for r in range(2, last + 1))


def test_data_headings_use_the_public_spelling(wb):
    assert wb["Data"]["J1"].value == "Backbone plasmid/gDNA"
    assert [wb["Data"][f"{c}1"].value for c in "MNO"] == [
        "PCR primer/Oligo 1", "PCR primer/Oligo 2", "Template"]


def test_hifi_re_formulas_look_up_only_the_four_column_table():
    ws = openpyxl.load_workbook(TEMPLATES["HiFi"])["RE"]
    refs, formulas = set(), 0
    for r in range(14, 26):
        for col in "FMN":
            value = ws[f"{col}{r}"].value
            formulas += 1
            refs.update(re.findall(r"RElist![$A-Z0-9:]+", value))
    assert formulas == 36
    assert refs == {"RElist!$A$4:$D$650", "RElist!$A$4:$A$650"}
    assert ws["C10"].value == "Yes"


def test_hifi_insert_count_survives_a_table_without_repeated_subs():
    ws = openpyxl.load_workbook(TEMPLATES["HiFi"])["HiFi"]
    assert ws["A8"].value == "=MAX(1,COUNTIF(L9:L20,MAX(1,IFERROR(MODE(L9:L20),1))))"


def test_hifi_blank_re2_stays_blank():
    ws = openpyxl.load_workbook(TEMPLATES["HiFi"])["RE"]
    assert ws["E14"].value == '=IF(Data!C2="","",IF(Data!I2="","",Data!I2))'


def check_relist(ws):
    """Rules every RElist must keep, bundled or edited: A-D filled, nothing outside, no dups."""
    names, problems = [], []
    for row in ws.iter_rows(min_row=4):
        values = [c.value for c in row]
        if all(v is None for v in values):
            continue
        a_to_d, rest = values[:4], values[4:]
        if any(v in (None, "") for v in a_to_d):
            problems.append((row[0].row, "A-D not all filled"))
        if any(v is not None for v in rest):
            problems.append((row[0].row, "value beyond column D"))
        names.append(str(a_to_d[0]))
        if "®" in str(a_to_d[0]):
            problems.append((row[0].row, "registered-trademark sign"))
    if len(names) != len(set(names)):
        problems.append((0, "duplicate enzyme names"))
    return problems


def test_bundled_relist_has_the_four_headings_and_no_enzymes():
    ws = openpyxl.load_workbook(TEMPLATES["HiFi"])["RElist"]
    assert [c.value for c in ws[3]] == ["Enzyme", "Buffer", "Reaction temp.",
                                        "Heat inactivation"]
    assert ws.auto_filter.ref is None
    assert check_relist(ws) == []
    assert all(c.value is None for row in ws.iter_rows(min_row=4) for c in row)
    assert ws["A1"].value and ws["A2"].value                  # the visible instructions


def test_relist_rules_hold_for_a_filled_table(enzyme_template):
    ws = openpyxl.load_workbook(enzyme_template)["RElist"]
    assert [ws.cell(row=r, column=1).value for r in range(4, 4 + len(ENZYMES))] == [
        e[0] for e in ENZYMES]
    assert check_relist(ws) == []


def test_relist_rules_catch_mistakes(enzyme_template):
    wb = openpyxl.load_workbook(enzyme_template)
    ws = wb["RElist"]
    ws["A20"], ws["B20"] = "Half®", "CutSmart"          # C, D empty; trademark sign
    ws["F21"] = "stray"
    ws["A22"] = "EcoRI-HF"                                    # duplicate
    for c, v in enumerate(["EcoRI-HF", "CutSmart", "37°C", "No"], 1):
        ws.cell(row=23, column=c, value=v)
    kinds = " ".join(msg for _, msg in check_relist(ws))
    assert "A-D not all filled" in kinds and "beyond column D" in kinds
    assert "registered-trademark" in kinds and "duplicate" in kinds


def notes_for(method):
    return [n for n in CONFIG["template_notes"] if n["template"] in (method, "both")]


def test_every_note_is_a_comment_in_the_template(wb, method):
    notes = notes_for(method)
    assert notes
    found = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.comment is not None:
                    found[(ws.title, c.coordinate)] = c.comment.text
    assert found == {(n["sheet"], n["cell"]): n["text"] for n in notes}


def test_golden_files_match_the_templates(method):
    stem = TEMPLATES[method].stem
    golden = (ROOT / "tests" / "golden" / f"{stem}.txt").read_text(encoding="utf-8")
    assert golden.splitlines() == dump_template.dump(TEMPLATES[method])


def test_dump_format(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "two\nlines\tand tab"
    ws["B1"] = "=A1&\"x\""
    ws["C1"] = 5
    ws["C1"].number_format = "00"
    path = tmp_path / "t.xlsx"
    wb.save(path)
    assert dump_template.dump(path) == [
        "S!A1\tvalue\ttwo\\nlines\\tand tab",
        'S!B1\tformula\t=A1&"x"',
        "S!C1\tvalue\t5",
        "S!C1\tnumfmt\t00",
    ]


def test_golden_files_are_utf8_with_the_degree_sign():
    text = Path(ROOT / "tests" / "golden" / "protocol.hifi.txt").read_bytes().decode("utf-8")
    assert "°C" in text and "×" in text


def test_the_columns_the_sheets_read_have_no_stray_formats():
    for method, path in TEMPLATES.items():
        ws = openpyxl.load_workbook(path)["Data"]
        last = CONFIG["methods"][method]["max_rows"] + 1
        for r in range(2, last + 1):
            for letter in "BGHIJLMNOP":
                assert ws[f"{letter}{r}"].number_format == "General", (method, letter, r)
